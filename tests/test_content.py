import asyncio
from contextlib import closing
import os
import random
import shutil
import sqlite3
import subprocess
import tempfile
import unittest
import wave
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch
from zipfile import ZipFile, ZIP_DEFLATED

import numpy as np
from docx import Document
from PIL import Image, ImageOps

from artifact_memory import Artifact, ArtifactMemory
from content_fusion import (AUDIO_METHODS, IMAGE_METHODS, donor_signal, fuse_audio,
                            fuse_bytes, fuse_image, fuse_text, merge_traces, trace)
from file_formats import ContentLimitError, ContentUnavailable, run_content_job, run_tool

_temp = tempfile.TemporaryDirectory()
with patch.dict(os.environ, {"BOT_TOKEN": "", "MEMORY_DB": str(Path(_temp.name) / "engine.sqlite3")}):
    import glitch_bot as engine


def image_bytes():
    yy, xx = np.indices((64, 96))
    pixels = np.stack((xx * 2, yy * 3, (xx + yy) % 255), axis=2).astype("uint8")
    output = BytesIO()
    Image.fromarray(pixels).save(output, "PNG")
    return output.getvalue()


def wav_bytes(seconds=1):
    t = np.arange(int(22050 * seconds)) / 22050
    samples = np.stack((np.sin(t * 440 * 2 * np.pi), np.sin(t * 660 * 2 * np.pi)), axis=1) * .35
    output = BytesIO()
    with wave.open(output, "wb") as target:
        target.setparams((2, 2, 22050, 0, "NONE", "not compressed"))
        target.writeframes((samples * 32767).astype("<i2").tobytes())
    return output.getvalue()


def docx_bytes():
    doc = Document()
    doc.add_heading("Original document voices", 0)
    p = doc.add_paragraph()
    p.add_run("These words come from the original document. ").bold = True
    p.add_run("Привет из исходного файла. Καλημέρα. " * 5)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Original table words and shared memories."
    table.cell(0, 1).text = "Another cell with some mathematical material."
    doc.sections[0].header.paragraphs[0].text = "An original header with several words."
    doc.add_picture(BytesIO(image_bytes()))
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def pdf_bytes(pages=2):
    """Small valid vector PDF with selectable text; no fixture download needed."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    kids = []
    font_index = 3 + pages * 2
    for i in range(pages):
        page_index, content_index = 3 + i * 2, 4 + i * 2
        kids.append(f"{page_index} 0 R")
        objects.append((f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 240 180] "
                        f"/Resources << /Font << /F1 {font_index} 0 R >> >> "
                        f"/Contents {content_index} 0 R >>").encode())
        text = f"BT /F1 14 Tf 12 140 Td (Original paper voices page {i + 1}) Tj ET".encode()
        objects.append(f"<< /Length {len(text)} >>\nstream\n".encode() + text + b"\nendstream")
    objects[1] = f"<< /Type /Pages /Count {pages} /Kids [{' '.join(kids)}] >>".encode()
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    output, offsets = bytearray(b"%PDF-1.4\n"), [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend((f"trailer << /Size {len(offsets)} /Root 1 0 R >>\n"
                   f"startxref\n{xref}\n%%EOF\n").encode())
    return bytes(output)


class MathTests(unittest.TestCase):
    def test_every_image_method_handles_small_and_rectangular_images(self):
        donor = trace(text="A different original donor with numbers 123456.", data=bytes(range(256)) * 10)
        for size in ((1, 1), (7, 11), (96, 64)):
            image = Image.new("RGB", size, (130, 50, 90))
            for method in IMAGE_METHODS:
                with self.subTest(method=method, size=size):
                    result, _ = fuse_image(image, trace(image=image), [donor], 12, [method])
                    self.assertEqual(result.size, size)
                    self.assertEqual(result.mode, "RGB")

    def test_actual_donor_values_change_image_and_audio_with_the_same_recipe(self):
        a = trace(waveform=np.linspace(-1, 1, 2048))
        b = trace(waveform=np.cos(np.linspace(0, 50, 2048)))
        image = Image.open(BytesIO(image_bytes())).convert("RGB")
        source = trace(image=image)
        outputs = []
        for donor in (a, b):
            random.seed(31)
            result, _ = fuse_image(image, source, [donor], 12, ["spectral crossover", "bitplane weave"])
            outputs.append(result.tobytes())
        self.assertNotEqual(*outputs)
        samples = np.sin(np.arange(4000) * .1).astype(np.float32)[:, None] * .3
        results = []
        for donor in (a, b):
            random.seed(31)
            results.append(fuse_audio(samples, trace(waveform=samples), [donor], 12,
                                      ["content convolution", "phase modulation"])[0])
        self.assertFalse(np.allclose(*results))

    def test_every_audio_method_preserves_length_and_finite_bounded_samples(self):
        donor = trace(text="A document can become a filter for an audio waveform.")
        for size in (1, 31, 1024, 1503, 22050):
            samples = np.repeat(np.sin(np.arange(size) * .13)[:, None], 2, axis=1).astype(np.float32) * .3
            for method in AUDIO_METHODS:
                with self.subTest(method=method, size=size):
                    output, _ = fuse_audio(samples, trace(waveform=samples), [donor], 12, [method])
                    self.assertEqual(output.shape, samples.shape)
                    self.assertTrue(np.isfinite(output).all())
                    self.assertLessEqual(float(np.max(np.abs(output))), .921)

    def test_whole_payload_and_decoded_material_contribute_to_bounded_traces(self):
        a = b"a" * 100000
        b = a[:70000] + b"z" * 10000 + a[80000:]
        first = trace(a, text="original words " * 1000, image=Image.new("RGB", (640, 480), "red"))
        second = trace(b, text="new words " * 1000, waveform=np.arange(10000))
        self.assertNotEqual(trace(a).signal, trace(b).signal)
        merged = merge_traces(first, second)
        self.assertLessEqual(len(merged.signal), 4096)
        self.assertLessEqual(len(merged.sample), 4096)
        self.assertLessEqual(len(merged.text), 5)
        self.assertLessEqual(max(Image.open(BytesIO(merged.image)).size), 320)
        self.assertTrue(merged.signal)

    def test_text_weave_uses_existing_unicode_words_and_separators(self):
        text = "One two\nПривет мир\tΚαλημέρα φίλε\n"
        output, labels = fuse_text(text, trace(text=text), [trace(text="other donor 123")], 12)
        self.assertCountEqual(text.split(), output.split())
        self.assertEqual([c for c in text if c.isspace()], [c for c in output if c.isspace()])
        self.assertTrue(labels)

    def test_algebraic_byte_fusion_retains_length_and_boundaries(self):
        data = bytes(range(256)) * 40
        for seed in range(10):
            random.seed(seed)
            result, recipe = fuse_bytes(data, trace(data), [trace(text="a document donor")], 12)
            self.assertEqual(len(result), len(data))
            self.assertEqual(result[:len(data) // 6], data[:len(data) // 6])
            self.assertEqual(result[len(data) * 5 // 6:], data[len(data) * 5 // 6:])
            self.assertNotEqual(result, data)
            self.assertTrue(recipe)


class MigrationTests(unittest.TestCase):
    def test_old_database_upgrades_and_preferences_do_not_reset_each_other(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "old.sqlite3"
            with closing(sqlite3.connect(path)) as db:
                db.executescript("""
                    CREATE TABLE artifacts (chat TEXT,digest TEXT,touched REAL,level INTEGER,
                        image BLOB,text TEXT,sample BLOB,PRIMARY KEY(chat,digest));
                    CREATE TABLE preferences (chat TEXT PRIMARY KEY,translate_enabled INTEGER,touched REAL);
                """)
                import time
                db.execute("INSERT INTO artifacts VALUES (?,?,?,1,?,?,?)", ("a", "x", time.time(), b"", '["old words"]', b"old"))
                db.commit()
            memory = ArtifactMemory(path)
            self.assertEqual(memory.snapshot("a", "x")[1][0].text, ("old words",))
            memory.translations("a", True)
            memory.settings("a", mode="raw", chaos=12)
            memory.translations("a", False)
            memory.translations("a", True)
            reopened = ArtifactMemory(path)
            self.assertEqual(reopened.settings("a"), ("raw", 12))
            self.assertTrue(reopened.translations("a"))
            artifact = trace(waveform=np.arange(9000))
            reopened.remember("a", "new", 12, artifact)
            self.assertTrue(reopened.snapshot("a", "new")[1][0].signal)
            self.assertEqual(reopened.settings("b"), ("fuse", 0))
            self.assertEqual(reopened.snapshot("b", "new")[1], [])
            reopened.forget("a")
            self.assertEqual(reopened.settings("a"), ("fuse", 0))
            self.assertFalse(reopened.translations("a"))
            self.assertEqual(reopened.snapshot("a", "")[1], [])


class FormatTests(unittest.IsolatedAsyncioTestCase):
    async def test_docx_is_reopenable_with_tables_runs_images_and_original_text_traces(self):
        data = docx_bytes()
        result, recipe, artifact, notices, name = await engine.process_upload(
            data, "paper.docx", 12, [trace(text="INHERITED donor phrase keeps returning.")])
        self.assertEqual(name, "paper.docx")
        self.assertIn("DOCX run weave", recipe)
        doc = Document(BytesIO(result))
        original = Document(BytesIO(data))
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertTrue(doc.paragraphs[1].runs[0].bold)
        self.assertNotEqual([p.text for p in doc.paragraphs], [p.text for p in original.paragraphs])
        with ZipFile(BytesIO(data)) as a, ZipFile(BytesIO(result)) as b:
            self.assertEqual(set(a.namelist()), set(b.namelist()))
            for name in a.namelist():
                if name.startswith("word/media/"):
                    self.assertNotEqual(a.read(name), b.read(name))
                    Image.open(BytesIO(b.read(name))).load()
        self.assertTrue(artifact.text)
        self.assertTrue(artifact.image)
        self.assertTrue(artifact.signal)
        self.assertTrue(notices)

    @unittest.skipUnless(shutil.which("pdfinfo") and shutil.which("pdftoppm"), "Poppler not installed")
    async def test_pdf_reopens_keeps_page_count_and_contributes_text(self):
        result, recipe, artifact, notices, name = await engine.process_upload(pdf_bytes(), "paper.pdf", 3, [])
        self.assertIn("PDF page raster", recipe)
        self.assertEqual(name, "paper.pdf")
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "result.pdf"
            path.write_bytes(result)
            info = subprocess.check_output(["pdfinfo", str(path)]).decode()
            self.assertRegex(info, r"Pages:\s+2")
        self.assertTrue(artifact.text)
        source_words = {"Original", "paper", "voices", "page", "1", "2"}
        self.assertTrue(all(set(snippet.split()) <= source_words for snippet in artifact.text))
        self.assertTrue(artifact.image)
        self.assertIn("flattened", " ".join(notices))

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "FFmpeg not installed")
    async def test_wav_reopens_preserves_sample_count_and_can_haunt_an_image(self):
        data = wav_bytes()
        result, recipe, artifact, notices, name = await engine.process_upload(
            data, "tone.wav", 12, [trace(text="A document becomes the room in which this sound is heard.")])
        self.assertEqual(name, "tone.wav")
        with wave.open(BytesIO(result)) as output, wave.open(BytesIO(data)) as original:
            self.assertEqual(output.getnframes(), original.getnframes())
            self.assertEqual(output.getframerate(), 22050)
            self.assertEqual(output.getnchannels(), 2)
            self.assertNotEqual(output.readframes(output.getnframes()), original.readframes(original.getnframes()))
        self.assertTrue(set(recipe) & set(AUDIO_METHODS))
        self.assertTrue(artifact.signal)
        random.seed(42)
        a, *_ = await engine.process_upload(image_bytes(), "after.png", 4, [artifact])
        random.seed(42)
        b, *_ = await engine.process_upload(image_bytes(), "after.png", 4, [])
        self.assertNotEqual(a, b)

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "FFmpeg not installed")
    async def test_mp4_retains_video_and_audio_and_decodes_all_the_way_through(self):
        with tempfile.TemporaryDirectory() as root:
            source = Path(root) / "source.mp4"
            subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-y", "-f", "lavfi", "-i",
                            "testsrc2=size=96x64:rate=8", "-f", "lavfi", "-i", "sine=frequency=440:sample_rate=22050",
                            "-t", "1", "-c:v", "libx264", "-threads", "1", "-pix_fmt", "yuv420p",
                            "-c:a", "aac", str(source)], check=True)
            for seed in (7, 42, 123):
                with self.subTest(seed=seed):
                    random.seed(seed)
                    result, recipe, artifact, notices, name = await engine.process_upload(
                        source.read_bytes(), "clip.mp4", 12, [trace(text="Video haunted by document words.")])
                    self.assertEqual(name, "clip.mp4")
                    self.assertIn("temporal feedback", recipe)
                    output = Path(root) / "result.mp4"
                    output.write_bytes(result)
                    subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-i", str(output),
                                    "-map", "0:v:0", "-map", "0:a:0", "-f", "null", "-"], check=True)
                    self.assertTrue(artifact.image)
                    self.assertTrue(artifact.signal)

    async def test_raw_mode_preserves_existing_byte_effects_without_decoder_calls(self):
        data = docx_bytes()
        with patch.object(engine, "run_content_job", side_effect=AssertionError("decoder called")):
            _, recipe, _, _, name = await engine.process_upload(data, "paper.docx", 4,
                                                               [Artifact(sample=b"donor" * 100)], mode="raw")
        self.assertEqual(recipe, ["inherited bytes", "byte chaos"])
        self.assertEqual(name, "paper.docx")

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "FFmpeg not installed")
    async def test_compressed_voice_audio_is_decoded_and_returned_with_a_wav_extension(self):
        with tempfile.TemporaryDirectory() as root:
            source = Path(root) / "tone.wav"
            source.write_bytes(wav_bytes(.4))
            for extension, codec in ((".ogg", "libopus"), (".mp3", "libmp3lame"), (".m4a", "aac")):
                with self.subTest(extension=extension):
                    target = source.with_suffix(extension)
                    subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-y", "-i", str(source),
                                    "-c:a", codec, str(target)], check=True)
                    result, recipe, _, _, name = await engine.process_upload(target.read_bytes(), "voice" + extension, 4, [])
                    self.assertEqual(name, "voice.wav")
                    self.assertTrue(set(recipe) & set(AUDIO_METHODS))
                    with wave.open(BytesIO(result)) as decoded:
                        self.assertGreater(decoded.getnframes(), 0)
                        self.assertEqual(decoded.getframerate(), 22050)

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "FFmpeg not installed")
    async def test_silent_gif_becomes_playable_mp4_without_an_audio_stream(self):
        first = Image.open(BytesIO(image_bytes())).convert("RGB")
        second = ImageOps.mirror(first)
        buf = BytesIO()
        first.save(buf, "GIF", save_all=True, append_images=[second], duration=250, loop=0)
        result, recipe, _, _, name = await engine.process_upload(buf.getvalue(), "animation.gif", 4, [])
        self.assertEqual(name, "animation.mp4")
        self.assertIn("temporal feedback", recipe)
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "result.mp4"
            path.write_bytes(result)
            subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-i", str(path), "-f", "null", "-"], check=True)

    async def test_decoder_failure_has_an_explicit_destructive_fallback(self):
        with patch.object(engine, "run_content_job", side_effect=ContentUnavailable("missing")):
            _, recipe, _, notices, _ = await engine.process_upload(b"broken pdf" * 30, "bad.pdf", 3, [])
        self.assertIn("byte chaos", recipe)
        self.assertIn("destructive", " ".join(notices))

    @unittest.skipUnless(shutil.which("pdfinfo"), "Poppler not installed")
    async def test_page_limit_is_reported_without_destructive_fallback(self):
        with self.assertRaisesRegex(ContentLimitError, "1–12"):
            await engine.process_upload(pdf_bytes(13), "long.pdf", 12, [])

    async def test_archive_expansion_is_rejected_before_xml_is_loaded(self):
        data = BytesIO()
        with ZipFile(data, "w", ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"x" * (33 * 1024 * 1024))
        with self.assertRaises(ContentLimitError):
            await engine.process_upload(data.getvalue(), "expanded.docx", 12, [])

    async def test_legacy_images_still_receive_old_effects_and_new_mathematics(self):
        _, recipe, _, _, _ = await engine.process_upload(image_bytes(), "a.png", 12, [])
        self.assertTrue(set(recipe) & set(engine.EFFECTS))
        self.assertTrue(set(recipe) & set(IMAGE_METHODS))


class WorkerDeadlineTests(unittest.TestCase):
    def test_decoder_cpu_limit_is_reported_as_a_limit(self):
        import signal
        with patch("file_formats.shutil.which", return_value="decoder"), patch(
                "file_formats.subprocess.run", side_effect=subprocess.CalledProcessError(-signal.SIGXCPU, "decoder")):
            with self.assertRaises(ContentLimitError):
                run_tool(["decoder"])

    def test_timeout_kills_the_decoder_process_group(self):
        process = Mock(pid=12345)
        process.wait.side_effect = [subprocess.TimeoutExpired("worker", 60), 0]
        with patch("file_formats.subprocess.Popen", return_value=process), patch("file_formats.os.killpg") as kill:
            with self.assertRaises(ContentLimitError):
                run_content_job(b"file", ".pdf", 1, [])
        kill.assert_called_once()
        self.assertEqual(kill.call_args.args[0], 12345)


class HandlerTests(unittest.IsolatedAsyncioTestCase):
    async def test_mode_and_chaos_commands_are_persistent_and_sender_scoped(self):
        message = SimpleNamespace(chat=SimpleNamespace(id=1), from_user=SimpleNamespace(id=2), answer=AsyncMock())
        with tempfile.TemporaryDirectory() as root:
            memory = ArtifactMemory(Path(root) / "memory.sqlite3")
            with patch.object(engine, "memory", memory):
                for command in ("/mode raw", "/chaos 12"):
                    message.text = command
                    await engine.handle_file(message)
                self.assertEqual(memory.settings("1:2"), ("raw", 12))
                self.assertEqual(memory.settings("1:3"), ("fuse", 0))
                message.text = "/chaos auto"
                await engine.handle_file(message)
                self.assertEqual(memory.settings("1:2"), ("raw", 0))
                message.text = "/chaos 99"
                await engine.handle_file(message)
                self.assertEqual(memory.settings("1:2"), ("raw", 0))
                message.text = "/forget"
                await engine.handle_file(message)
                self.assertEqual(memory.settings("1:2"), ("fuse", 0))

    async def test_voice_upload_and_sender_settings_reach_the_processor(self):
        upload = SimpleNamespace(file_size=128, file_id="voice", file_name=None)
        message = SimpleNamespace(chat=SimpleNamespace(id=1), from_user=SimpleNamespace(id=2),
                                  text=None, photo=[], video=None, audio=None, voice=upload,
                                  video_note=None, animation=None, document=None,
                                  answer=AsyncMock(), reply_document=AsyncMock())
        bot = SimpleNamespace(get_file=AsyncMock(return_value=SimpleNamespace(file_path="voice.ogg")),
                              download_file=AsyncMock(return_value=BytesIO(b"uploaded bytes")))
        with tempfile.TemporaryDirectory() as root:
            memory = ArtifactMemory(Path(root) / "memory.sqlite3")
            memory.settings("1:2", mode="raw", chaos=12)
            processor = AsyncMock(return_value=(b"result", ["byte chaos"], Artifact(sample=b"trace"), [], "voice.ogg"))
            with patch.object(engine, "memory", memory), patch.object(engine, "bot", bot), patch.object(engine, "process_upload", processor):
                await engine.handle_file(message)
            self.assertEqual(processor.call_args.args[1:3], ("voice.ogg", 12))
            self.assertEqual(processor.call_args.args[-1], "raw")
            message.reply_document.assert_awaited_once()
            self.assertEqual(len(memory.snapshot("1:2", "")[1]), 1)
            self.assertEqual(memory.snapshot("1:3", "")[1], [])

    async def test_failed_delivery_is_never_remembered(self):
        upload = SimpleNamespace(file_size=128, file_id="file", file_name="paper.pdf")
        message = SimpleNamespace(chat=SimpleNamespace(id=1), from_user=SimpleNamespace(id=2),
                                  text=None, photo=[], video=None, audio=None, voice=None,
                                  video_note=None, animation=None, document=upload,
                                  answer=AsyncMock(), reply_document=AsyncMock(side_effect=RuntimeError("delivery")))
        bot = SimpleNamespace(get_file=AsyncMock(return_value=SimpleNamespace(file_path="paper.pdf")),
                              download_file=AsyncMock(return_value=BytesIO(b"uploaded bytes")))
        with tempfile.TemporaryDirectory() as root:
            memory = ArtifactMemory(Path(root) / "memory.sqlite3")
            with patch.object(engine, "memory", memory), patch.object(engine, "bot", bot), patch.object(
                    engine, "process_upload", AsyncMock(return_value=(b"out", [], Artifact(sample=b"trace"), [], "paper.pdf"))):
                await engine.handle_file(message)
            self.assertEqual(memory.snapshot("1:2", "")[1], [])
            message.answer.assert_awaited_once()
